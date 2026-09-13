#!/usr/bin/env python3
"""Build a branded Eventa PRD Word document from docs/PRD_EVENTA_SAAS.md."""

from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PRD_EVENTA_SAAS.md"
OUTPUT = ROOT / "docs" / "Eventa_SaaS_PRD_v1.0.docx"
ARTIFACT = Path("/opt/cursor/artifacts/Eventa_SaaS_PRD_v1.0.docx")

ESPRESSO = RGBColor(0x1A, 0x0F, 0x0A)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GOLD_MUTED = RGBColor(0x8A, 0x6F, 0x38)
BODY = RGBColor(0x3A, 0x24, 0x19)
CREAM = "FAF6EF"


def set_run_font(run, name: str = "Calibri", size: Pt | None = None, bold: bool | None = None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_paragraph(paragraph, fill_hex: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    ppr.append(shd)


def set_cell_shading(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    heading_sizes = {1: 20, 2: 15, 3: 13}
    for level, size in heading_sizes.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = ESPRESSO if level == 1 else GOLD_MUTED
        style.paragraph_format.space_before = Pt(16 if level == 1 else 12)
        style.paragraph_format.space_after = Pt(8)

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(32)
    title.font.bold = True
    title.font.color.rgb = ESPRESSO

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(16)
    subtitle.font.color.rgb = GOLD_MUTED


def build_reference_doc(path: Path) -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    doc.save(path)


def insert_title_page(doc: Document) -> None:
    body = doc.element.body
    first = body[0]

    def add_para(text: str, *, size=12, bold=False, color=BODY, align="left", space_before=0, space_after=6, shade=None):
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), align)
        ppr.append(jc)
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:before"), str(int(space_before * 20)))
        sp.set(qn("w:after"), str(int(space_after * 20)))
        ppr.append(sp)
        if shade:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), shade)
            ppr.append(shd)
        p.append(ppr)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        rf.set(qn("w:ascii"), "Calibri")
        rf.set(qn("w:hAnsi"), "Calibri")
        rpr.append(rf)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(size * 2))
        rpr.append(sz)
        if bold:
            rpr.append(OxmlElement("w:b"))
        c = OxmlElement("w:color")
        c.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
        rpr.append(c)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        p.append(r)
        first.addprevious(p)

    add_para("PROMAX IT SOLUTIONS", size=12, bold=True, color=GOLD_MUTED, space_before=72, space_after=18)
    add_para("EVENTA", size=36, bold=True, color=ESPRESSO, space_after=6)
    add_para("Product Requirements Document", size=20, color=GOLD_MUTED, space_after=18)
    add_para("Multi-tenant event management SaaS", size=14, color=BODY, space_after=28)
    add_para("Commercial name: Eventa by Promax", size=12, color=BODY, space_after=4)
    add_para("Version 1.0  ·  Draft for review  ·  13 September 2026", size=12, color=BODY, space_after=4)
    add_para("Reference tenant: Yoruba Day Canberra 2026", size=12, color=BODY, space_after=4)
    add_para("Target stack: React  ·  ASP.NET Core  ·  Microsoft SQL Server", size=12, color=BODY, space_after=28)
    add_para(
        "This document generalises the production Promax Event Platform into a commercial "
        "SaaS product. Yoruba Day Canberra remains tenant zero — a configuration, not a code fork.",
        size=11,
        color=BODY,
        space_after=36,
    )
    add_para("Classification: Internal — Promax IT Solutions and authorised partners", size=10, color=GOLD_MUTED, space_after=0)

    br = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    pb = OxmlElement("w:pageBreakBefore")
    # Use an explicit break so the PRD starts on page 2.
    ppr.append(OxmlElement("w:spacing"))
    br.append(ppr)
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:br"))
    r[-1].set(qn("w:type"), "page")
    br.append(r)
    first.addprevious(br)


def apply_headers_footers(doc: Document) -> None:
    for section in doc.sections:
        section.different_first_page_header_footer = True
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.clear()
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = hp.add_run("Eventa by Promax")
        set_run_font(run, size=Pt(9), bold=True, color=ESPRESSO)
        run = hp.add_run("   ·   Product Requirements Document  ·  v1.0")
        set_run_font(run, size=Pt(9), color=GOLD_MUTED)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = fp.add_run("© 2026 Promax IT Solutions  ·  Confidential  ·  Page ")
        set_run_font(run, size=Pt(9), color=GOLD_MUTED)
        add_page_number(fp)
        run = fp.add_run("  ·  eventa.app")
        set_run_font(run, size=Pt(9), color=GOLD_MUTED)

        # First page: lighter header
        fh = section.first_page_header
        fh.is_linked_to_previous = False
        fhp = fh.paragraphs[0]
        fhp.clear()
        fhp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = fhp.add_run("Promax IT Solutions")
        set_run_font(run, size=Pt(9), bold=True, color=GOLD_MUTED)

        ff = section.first_page_footer
        ff.is_linked_to_previous = False
        ffp = ff.paragraphs[0]
        ffp.clear()
        ffp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = ffp.add_run("Draft for review  ·  13 September 2026")
        set_run_font(run, size=Pt(9), color=GOLD_MUTED)


def style_tables(doc: Document) -> None:
    for table in doc.tables:
        table.autofit = True
        tbl = table._tbl
        tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "E8DFD0")
            borders.append(el)
        tbl_pr.append(borders)
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                fill = "1A0F0A" if i == 0 else ("FAF6EF" if i % 2 == 0 else "FFFFFF")
                set_cell_shading(cell, fill)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.name = "Calibri"
                        run.font.size = Pt(9)
                        if i == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(0xFA, 0xF6, 0xEF)
                        else:
                            run.font.color.rgb = BODY


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "Eventa — Product Requirements Document"
    props.author = "Promax IT Solutions"
    props.subject = "Multi-tenant event management SaaS PRD"
    props.category = "Product Requirements"
    props.comments = "Eventa by Promax. Reference tenant: Yoruba Day Canberra 2026."
    props.keywords = "Eventa, Promax, SaaS, PRD, multi-tenant, events"


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source: {SOURCE}")

    with tempfile.TemporaryDirectory() as tmp:
        ref = Path(tmp) / "reference.docx"
        raw = Path(tmp) / "raw.docx"
        build_reference_doc(ref)
        subprocess.run(
            [
                "pandoc",
                str(SOURCE),
                "--from",
                "markdown+pipe_tables+grid_tables+fenced_code_blocks+auto_identifiers",
                "--to",
                "docx",
                "--reference-doc",
                str(ref),
                "--toc",
                "--toc-depth=2",
                "-o",
                str(raw),
            ],
            check=True,
        )
        doc = Document(raw)
        configure_styles(doc)
        insert_title_page(doc)
        apply_headers_footers(doc)
        style_tables(doc)
        set_core_properties(doc)
        OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        doc.save(OUTPUT)

    ARTIFACT.parent.mkdir(parents=True, exist_ok=True)
    ARTIFACT.write_bytes(OUTPUT.read_bytes())
    print(f"Wrote {OUTPUT} ({OUTPUT.stat().st_size} bytes)")
    print(f"Wrote {ARTIFACT} ({ARTIFACT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
